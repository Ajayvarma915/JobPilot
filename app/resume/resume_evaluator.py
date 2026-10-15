from __future__ import annotations

import re
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader


STANDARD_HEADINGS = {
    "technical skills",
    "skills",
    "experience",
    "work experience",
    "projects",
    "education",
    "certifications",
    "certificates",
    "achievements",
    "summary",
    "professional summary",
}
ACTION_VERBS = {
    "built", "created", "developed", "designed", "implemented", "integrated",
    "deployed", "automated", "optimized", "improved", "reduced", "increased",
    "delivered", "engineered", "configured", "tested", "analyzed", "led",
    "migrated", "refactored", "launched", "secured", "maintained", "supported",
    "constructed", "established", "streamlined", "resolved", "implemented",
}
RESULT_TERMS = {
    "improved", "reduced", "increased", "accelerated", "saved", "achieved",
    "delivered", "enabled", "supported", "automated", "resolved", "optimized",
    "accuracy", "latency", "uptime", "users", "requests", "tests", "tasks",
}
TECH_TERMS = {
    "python", "java", "javascript", "typescript", "react", "next.js", "node.js",
    "fastapi", "sql", "sqlite", "postgresql", "mongodb", "docker", "aws",
    "azure", "gcp", "git", "rest api", "api", "auth.js", "firebase", "html",
    "css", "jest", "pytest", "linux", "kubernetes", "redis", "machine learning",
}


def _normal(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip().casefold()


def _contains_phrase(text: str, phrase: str) -> bool:
    escaped = re.escape(phrase.strip())
    return bool(re.search(rf"(?<![a-z0-9+#]){escaped}(?![a-z0-9+#])", text, re.IGNORECASE))


def _docx_payload(path: str | Path) -> dict[str, Any]:
    document = Document(str(path))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
    table_text = [cell.text.strip() for table in document.tables for row in table.rows for cell in row.cells if cell.text.strip()]
    text = "\n".join([*paragraphs, *table_text])
    sizes: list[float] = []
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if run.text.strip() and run.font.size is not None:
                sizes.append(float(run.font.size.pt))
    multi_column_sections = 0
    for section in document.sections:
        try:
            cols = section._sectPr.xpath("./w:cols")
            if cols and cols[0].get("num") and int(cols[0].get("num")) > 1:
                multi_column_sections += 1
        except (ValueError, TypeError):
            continue
    return {
        "text": text,
        "paragraphs": paragraphs,
        "table_count": len(document.tables),
        "image_count": len(document.inline_shapes),
        "font_sizes": sizes,
        "multi_column_sections": multi_column_sections,
        "page_width_inches": float(document.sections[0].page_width.inches) if document.sections else None,
        "page_height_inches": float(document.sections[0].page_height.inches) if document.sections else None,
    }


def _pdf_payload(path: str | Path) -> dict[str, Any]:
    reader = PdfReader(str(path))
    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return {"text": "\n".join(pages), "page_count": len(reader.pages)}


def _heading_count(text: str) -> int:
    lines = {_normal(line).strip(" :-") for line in text.splitlines() if line.strip()}
    normalized = _normal(text)
    # Exact-line match is the main signal. Contained match is a fallback for
    # parsers that flatten line breaks while preserving section text.
    count = sum(1 for heading in STANDARD_HEADINGS if heading in lines)
    if count < 3:
        count = max(count, sum(1 for heading in STANDARD_HEADINGS if re.search(rf"\b{re.escape(heading)}\b", normalized)))
    return count


def _score_label(score: int, *, ats: bool = False) -> str:
    if ats:
        if score >= 90:
            return "Strong parsing signals"
        if score >= 75:
            return "Generally readable; review flags"
        return "Parsing risks detected"
    if score >= 90:
        return "Strong"
    if score >= 75:
        return "Good foundation"
    if score >= 55:
        return "Needs improvement"
    return "Limited evidence"


def _ats_readability(docx_path: str | Path, pdf_path: str | Path, candidate_name: str) -> dict[str, Any]:
    checks: list[dict[str, Any]] = []
    issues: list[str] = []
    docx = _docx_payload(docx_path)
    pdf = _pdf_payload(pdf_path)
    docx_text = docx["text"]
    pdf_text = pdf["text"]
    name = _normal(candidate_name)

    def add(label: str, passed: bool, points: int, detail: str, problem: str | None = None) -> None:
        checks.append({"label": label, "passed": bool(passed), "points": points if passed else 0, "detail": detail})
        if not passed and problem:
            issues.append(problem)

    both_extracted = len(docx_text.strip()) >= 80 and len(pdf_text.strip()) >= 80 and pdf["page_count"] >= 1
    add(
        "Selectable text in DOCX and PDF", both_extracted, 20,
        f"DOCX extracted {len(docx_text.strip())} characters; PDF extracted {len(pdf_text.strip())} characters across {pdf['page_count']} page(s).",
        "One or both formats did not yield enough selectable text. Recheck the generated files before applying.",
    )

    name_found = bool(name) and name in _normal(docx_text) and name in _normal(pdf_text)
    add("Candidate name is extractable", name_found, 15,
        "The candidate name was found in both generated files." if name_found else "The candidate name was not found in both files.",
        "Confirm the candidate name is included and extractable in both files.")

    docx_headings = _heading_count(docx_text)
    pdf_headings = _heading_count(pdf_text)
    headings_found = min(docx_headings, pdf_headings) >= 3
    add("Conventional section headings", headings_found, 20,
        f"Detected {docx_headings} recognized headings in DOCX and {pdf_headings} in PDF.",
        "Use standard headings such as Technical Skills, Experience, Projects, Education, and Certifications.")

    no_complex_layout = (
        docx["table_count"] == 0
        and docx["image_count"] == 0
        and docx["multi_column_sections"] == 0
    )
    add("Single-column text-first layout", no_complex_layout, 15,
        f"DOCX contains {docx['table_count']} table(s), {docx['image_count']} inline image(s), and {docx['multi_column_sections']} multi-column section(s).",
        "Prefer a one-column text-first layout; remove tables, photos, and text boxes when they are not required.")

    sizes = docx["font_sizes"]
    body_sizes = [size for size in sizes if 8 <= size <= 12]
    font_readable = bool(sizes) and len(body_sizes) / len(sizes) >= 0.8 and max(sizes, default=0) <= 24
    add("Readable, consistent font sizes", font_readable, 10,
        f"{len(body_sizes)} of {len(sizes)} explicitly sized text runs use 8–12 pt; document title/heading runs may be larger.",
        "Review typography: keep body text roughly 9–11 pt and section headings consistent.")

    contact = bool(re.search(r"\b[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\b", docx_text)) and bool(
        re.search(r"\b[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\b", pdf_text)
    )
    add("Contact email is extractable", contact, 10,
        "An email address was found in both formats." if contact else "No email address was found in both formats.",
        "Verify that a correct contact email appears near the top of the résumé.")

    words = len(re.findall(r"\b[\w+#.-]+\b", pdf_text))
    pages_ok = 1 <= pdf["page_count"] <= 2
    useful_text = words >= 100
    volume_ok = pages_ok and useful_text
    add("Reasonable length and extracted content", volume_ok, 10,
        f"PDF has {pdf['page_count']} page(s) and about {words} extractable words.",
        "Review page count and content density. Avoid shrinking text or adding filler just to fill a page.")

    score = min(100, max(0, sum(int(item["points"]) for item in checks)))
    return {
        "score": score,
        "label": _score_label(score, ats=True),
        "checks": checks,
        "issues": issues,
        "pdf_page_count": pdf["page_count"],
        "pdf_extractable_word_count": words,
        "disclaimer": "This is JobPilot's local file-readability heuristic, not a score assigned by Google, Amazon, or any employer's ATS.",
    }


def _coverage_score(items: Iterable[str], covered: Iterable[str]) -> tuple[int | None, list[str], list[str]]:
    wanted = list(dict.fromkeys(str(item).strip() for item in items if str(item).strip()))
    matched = {str(item).casefold() for item in covered}
    if not wanted:
        return None, [], []
    matched_display = [item for item in wanted if item.casefold() in matched]
    missing = [item for item in wanted if item.casefold() not in matched]
    return round(len(matched_display) / len(wanted) * 100), matched_display, missing


def _job_relevance(selection: dict[str, Any]) -> dict[str, Any]:
    signals = selection.get("signals", {}) if isinstance(selection.get("signals"), dict) else {}
    coverage = selection.get("coverage", {}) if isinstance(selection.get("coverage"), dict) else {}
    selected = selection.get("selected_evidence", [])
    if not isinstance(selected, list):
        selected = []

    required = list(signals.get("required_skills", []) or [])
    preferred = list(signals.get("preferred_skills", []) or [])
    inferred = list(signals.get("inferred_job_technologies", []) or [])
    required_covered = list(coverage.get("verified_covered", []) or [])
    required_pending = list(coverage.get("awaiting_evidence_review", []) or [])
    required_missing = list(coverage.get("not_found_in_evidence", []) or [])
    required_score, required_matched, required_missing_calc = _coverage_score(required, required_covered)

    matched_preferred = list(dict.fromkeys(
        skill for item in selected for skill in (item.get("matched_preferred_skills", []) or [])
    ))
    preferred_score, preferred_matched, preferred_missing = _coverage_score(preferred, matched_preferred)
    matched_inferred = list(dict.fromkeys(
        skill for item in selected for skill in (item.get("matched_job_technologies", []) or [])
    ))
    inferred_score, inferred_matched_display, inferred_missing = _coverage_score(inferred, matched_inferred)

    categories = {str(item.get("category", "")).casefold() for item in selected}
    depth_score = 0
    if "project" in categories:
        depth_score += 50
    if "experience" in categories:
        depth_score += 35
    if "skill" in categories:
        depth_score += 15
    if "education" in categories:
        depth_score += 5
    depth_score = min(100, depth_score)

    components: list[tuple[str, float, float]] = []
    if required_score is not None:
        components.append(("Required skill coverage", required_score, 0.60))
    if preferred_score is not None:
        components.append(("Preferred skill coverage", preferred_score, 0.15 if required_score is not None else 0.35))
    if inferred_score is not None:
        components.append(("JD technology coverage", inferred_score, 0.15 if required_score is not None else (0.30 if preferred_score is not None else 0.40)))
    assigned_weight = sum(weight for _, _, weight in components)
    depth_weight = max(0.0, 1.0 - assigned_weight)
    components.append(("Relevant verified evidence", depth_score, depth_weight))
    total_weight = sum(weight for _, _, weight in components)
    score = round(sum(value * weight for _, value, weight in components) / total_weight) if total_weight else depth_score

    blockers: list[str] = []
    if required_missing or required_missing_calc:
        blockers.append("One or more required skills have no matching verified evidence.")
    if required_pending:
        blockers.append("Some required skills appear only in evidence awaiting review.")
    if required_score is None:
        blockers.append("No explicit required skills were supplied; relevance is based on available job signals and selected evidence.")

    return {
        "score": max(0, min(100, score)),
        "label": _score_label(max(0, min(100, score))),
        "required_skill_coverage": {
            "score": required_score,
            "matched": required_matched,
            "missing": list(dict.fromkeys(required_missing or required_missing_calc)),
            "awaiting_review": required_pending,
        },
        "preferred_skill_coverage": {
            "score": preferred_score,
            "matched": preferred_matched,
            "missing": preferred_missing,
        },
        "job_technology_coverage": {
            "score": inferred_score,
            "matched": inferred_matched_display,
            "missing": inferred_missing,
        },
        "verified_evidence_categories": sorted(categories),
        "blockers": blockers,
        "disclaimer": "This is JobPilot's evidence-based heuristic; it is not an employer's ranking or a prediction of shortlisting.",
    }


def _bullet_text(item: dict[str, Any]) -> list[str]:
    metadata = item.get("metadata", {})
    bullets = metadata.get("bullets", []) if isinstance(metadata, dict) else []
    if isinstance(bullets, list) and bullets:
        return [str(value).strip() for value in bullets if str(value).strip()]
    statement = str(item.get("statement", "")).strip()
    return [statement] if statement else []


def _score_bullet(text: str, technologies: set[str]) -> dict[str, Any]:
    normalized = _normal(text)
    words = re.findall(r"[a-zA-Z0-9+#.-]+", text)
    first = words[0].casefold().strip(".,:;()[]{}") if words else ""
    action = first in ACTION_VERBS
    tech_detail = any(_contains_phrase(normalized, term) for term in TECH_TERMS | technologies)
    concrete_outcome = any(_contains_phrase(normalized, term) for term in RESULT_TERMS) or bool(re.search(r"\d", text))
    useful_length = 10 <= len(words) <= 38
    specific = len(set(word.casefold() for word in words if len(word) >= 4)) >= 7 and not re.search(
        r"\b(responsible for|worked on various|helped with|hardworking|team player)\b", normalized
    )
    points = 20 * sum([action, tech_detail, concrete_outcome, useful_length, specific])
    return {
        "score": points,
        "signals": {
            "starts_with_action_verb": action,
            "names_technology_or_specific_system": tech_detail,
            "includes_outcome_or_measurable_scope": concrete_outcome,
            "has_concise_detail": useful_length,
            "is_specific_not_generic": specific,
        },
    }


def _achievement_quality(selection: dict[str, Any]) -> dict[str, Any]:
    selected = selection.get("selected_evidence", [])
    if not isinstance(selected, list):
        selected = []
    focus = [item for item in selected if str(item.get("category", "")).casefold() in {"project", "experience", "achievement"}]
    bullets: list[dict[str, Any]] = []
    technologies: set[str] = set()
    for item in focus:
        metadata = item.get("metadata", {})
        if isinstance(metadata, dict):
            technologies.update(str(value).casefold() for value in (metadata.get("technologies", []) or []))
        for text in _bullet_text(item):
            bullets.append({"text": text, **_score_bullet(text, technologies)})

    if not bullets:
        return {
            "score": 0,
            "label": "Insufficient detail to assess",
            "sample_size": 0,
            "checks": [],
            "recommendation": "Add verified project or experience details to the master résumé before tailoring this role.",
        }

    score = round(sum(item["score"] for item in bullets) / len(bullets))
    signals = [item["signals"] for item in bullets]
    checks = []
    for key, label in (
        ("starts_with_action_verb", "Starts with an action verb"),
        ("names_technology_or_specific_system", "Names a technology or specific system"),
        ("includes_outcome_or_measurable_scope", "Provides an outcome or measurable scope"),
        ("has_concise_detail", "Provides concise, useful detail"),
        ("is_specific_not_generic", "Avoids generic wording"),
    ):
        passing = sum(1 for item in signals if item[key])
        checks.append({"label": label, "passed_count": passing, "sample_size": len(signals)})

    if score >= 90:
        recommendation = "The selected bullets show strong detail. Verify every result and metric before applying."
    elif score >= 75:
        recommendation = "Strengthen bullets with a concrete result or verified scope where available. Never invent metrics."
    else:
        recommendation = "Review the source evidence and improve vague bullets with the action, technical approach, and a verifiable outcome where available."
    return {
        "score": score,
        "label": _score_label(score),
        "sample_size": len(bullets),
        "checks": checks,
        "recommendation": recommendation,
    }


def evaluate_generated_resume(
    docx_path: str | Path,
    pdf_path: str | Path,
    selection: dict[str, Any],
    store: dict[str, Any],
) -> dict[str, Any]:
    """Return transparent local checks for generated DOCX/PDF and job fit.

    Scores describe file readability and source-evidence relevance. They do not
    emulate a vendor ATS model and must never be represented as a shortlist
    probability or employer-assigned score.
    """
    candidate = store.get("candidate", {}) if isinstance(store.get("candidate"), dict) else {}
    name = str(candidate.get("name", "")).strip()
    ats = _ats_readability(docx_path, pdf_path, name)
    relevance = _job_relevance(selection)
    quality = _achievement_quality(selection)

    recommendations: list[str] = []
    recommendations.extend(ats.get("issues", []))
    recommendations.extend(relevance.get("blockers", []))
    if quality.get("recommendation"):
        recommendations.append(str(quality["recommendation"]))
    if selection.get("review_queue"):
        recommendations.append("Review relevant unverified evidence in the evidence store; it was not included in the generated résumé.")
    if not recommendations:
        recommendations.append("Core checks passed. Review both files carefully and confirm every detail before applying.")

    missing = relevance.get("required_skill_coverage", {}).get("missing", [])
    pending = relevance.get("required_skill_coverage", {}).get("awaiting_review", [])
    ready = ats["score"] >= 90 and relevance["score"] >= 80 and quality["score"] >= 70 and not missing and not pending
    return {
        "readiness_label": "Ready for your review" if ready else "Needs review before applying",
        "readiness_passes_local_checks": bool(ready),
        "ats_readability": ats,
        "job_relevance": relevance,
        "achievement_quality": quality,
        "recommendations": list(dict.fromkeys(recommendations)),
        "global_disclaimer": "These are JobPilot's local heuristic checks, not official Google/Amazon ATS scores, selection criteria, or a guarantee of interviews. Always follow the employer's file-format instructions.",
    }


__all__ = ["evaluate_generated_resume"]
