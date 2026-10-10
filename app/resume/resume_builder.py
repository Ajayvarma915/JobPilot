from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


class ResumeBuildError(ValueError):
    """Raised when a selection is unsafe, stale, or invalid for generation."""


NAVY = RGBColor(31, 55, 86)
DARK = RGBColor(35, 38, 42)
MUTED = RGBColor(86, 94, 105)


def _read_json(path: str | Path) -> dict[str, Any]:
    resolved = Path(path).expanduser().resolve()

    if not resolved.exists():
        raise FileNotFoundError(f"File not found: {resolved}")

    try:
        value = json.loads(resolved.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ResumeBuildError(f"Invalid JSON file: {resolved}") from exc

    if not isinstance(value, dict):
        raise ResumeBuildError(f"Expected a JSON object in {resolved}")

    return value


def _safe_slug(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "_", value).strip("_")
    return slug[:80] or "Tailored_Resume"


def _utc_now() -> str:
    return (
        datetime.now(timezone.utc)
        .isoformat(timespec="seconds")
        .replace("+00:00", "Z")
    )


def _set_run_font(
    run,
    *,
    size: float | None = None,
    bold: bool | None = None,
    color: RGBColor | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")

    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.italic = italic


def _set_paragraph_spacing(
    paragraph,
    *,
    before: float = 0,
    after: float = 2,
    line_spacing: float = 1.0,
) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line_spacing
    paragraph.paragraph_format.keep_together = True


def _add_text(
    document: Document,
    text: str,
    *,
    size: float = 9.5,
    bold: bool = False,
    color: RGBColor = DARK,
    italic: bool = False,
    align=None,
    after: float = 2,
):
    paragraph = document.add_paragraph()

    if align is not None:
        paragraph.alignment = align

    _set_paragraph_spacing(paragraph, after=after)
    run = paragraph.add_run(text)

    _set_run_font(
        run,
        size=size,
        bold=bold,
        color=color,
        italic=italic,
    )

    return paragraph


def _add_bullet(document: Document, text: str) -> None:
    cleaned = " ".join(str(text).split()).strip()

    if not cleaned:
        return

    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.first_line_indent = Inches(-0.14)

    _set_paragraph_spacing(
        paragraph,
        after=1,
        line_spacing=1.0,
    )

    run = paragraph.add_run(cleaned)
    _set_run_font(run, size=9.1, color=DARK)


def _add_section_heading(document: Document, title: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True

    _set_paragraph_spacing(
        paragraph,
        before=5,
        after=2,
    )

    run = paragraph.add_run(title.upper())
    _set_run_font(
        run,
        size=10.5,
        bold=True,
        color=NAVY,
    )

    # Add a subtle bottom border without using a layout table.
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")

    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "5")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "C6D0DB")

    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _add_entry_heading(
    document: Document,
    title: str,
    detail: str | None = None,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True

    _set_paragraph_spacing(
        paragraph,
        before=2,
        after=1,
    )

    run = paragraph.add_run(" ".join(str(title).split()))
    _set_run_font(
        run,
        size=9.6,
        bold=True,
        color=DARK,
    )

    if detail:
        detail_run = paragraph.add_run(
            f"  |  {' '.join(str(detail).split())}"
        )
        _set_run_font(
            detail_run,
            size=8.9,
            color=MUTED,
        )


def _validated_records(
    selection: dict[str, Any],
    store: dict[str, Any],
) -> list[dict[str, Any]]:
    safety = selection.get("safety", {})

    if safety.get("generated_claims") is True:
        raise ResumeBuildError(
            "Selection report indicates generated claims; refusing to build."
        )

    if safety.get("only_user_verified_records_selected") is not True:
        raise ResumeBuildError(
            "Selection report lacks the required verification-safety flag."
        )

    selected = selection.get("selected_evidence")

    if not isinstance(selected, list) or not selected:
        raise ResumeBuildError(
            "No verified evidence was selected. Review evidence or select a suitable JD first."
        )

    current_by_id = {
        item.get("evidence_id"): item
        for item in store.get("evidence", [])
        if item.get("evidence_id")
    }

    validated: list[dict[str, Any]] = []
    seen: set[str] = set()

    for selected_item in selected:
        evidence_id = selected_item.get("evidence_id")

        if not evidence_id or evidence_id in seen:
            raise ResumeBuildError(
                "Selection contains a missing or duplicate evidence ID."
            )

        seen.add(evidence_id)
        current = current_by_id.get(evidence_id)

        if current is None:
            raise ResumeBuildError(
                f"Selected evidence no longer exists in the store: {evidence_id}"
            )

        if current.get("review", {}).get("user_verified") is not True:
            raise ResumeBuildError(
                f"Selected evidence is not currently verified: "
                f"{current.get('title', evidence_id)}. Re-run selection."
            )

        if selected_item.get("user_verified") is not True:
            raise ResumeBuildError(
                f"Selection report does not mark evidence as verified: "
                f"{current.get('title', evidence_id)}"
            )

        # Reject a stale report when the evidence has changed.
        selected_source = selected_item.get("source", {})
        current_source = current.get("source", {})

        if selected_item.get("statement", "") != current.get("statement", ""):
            raise ResumeBuildError(
                f"Selection report is stale for "
                f"'{current.get('title', evidence_id)}'; re-run evidence selection."
            )

        if selected_source.get("text", "") != current_source.get("text", ""):
            raise ResumeBuildError(
                f"Source evidence changed for "
                f"'{current.get('title', evidence_id)}'; re-run evidence selection."
            )

        # The current store is authoritative. The report provides ranking only.
        merged = dict(current)
        merged["score"] = selected_item.get("score", 0)
        merged["reason"] = selected_item.get("reason", [])
        merged["matched_required_skills"] = selected_item.get(
            "matched_required_skills", []
        )
        merged["matched_preferred_skills"] = selected_item.get(
            "matched_preferred_skills", []
        )

        validated.append(merged)

    # Education is a core resume section. Include currently verified education
    # even when the job-description selector found no keyword overlap.
    selected_ids = {
        item.get("evidence_id")
        for item in validated
    }

    for current in store.get("evidence", []):
        if current.get("category") != "education":
            continue

        if current.get("review", {}).get("user_verified") is not True:
            continue

        evidence_id = current.get("evidence_id")

        if not evidence_id or evidence_id in selected_ids:
            continue

        baseline = dict(current)
        baseline["score"] = 0
        baseline["reason"] = [
            "included as a core resume section: verified education"
        ]
        baseline["included_as_core_section"] = True

        validated.append(baseline)
        selected_ids.add(evidence_id)

    return validated


def _contact_lines(candidate: dict[str, Any]) -> list[str]:
    parts = [
        candidate.get("email"),
        candidate.get("phone"),
        candidate.get("location"),
    ]

    line1 = "  |  ".join(
        str(item).strip()
        for item in parts
        if item and str(item).strip()
    )

    links = [
        candidate.get("linkedin"),
        candidate.get("github"),
    ]

    line2 = "  |  ".join(
        str(item).strip()
        for item in links
        if item and str(item).strip()
    )

    return [
        line
        for line in (line1, line2)
        if line
    ]


def _skills_text(records: list[dict[str, Any]]) -> str:
    # Use only the canonical title of each selected, verified skill record.
    return ", ".join(
        dict.fromkeys(
            str(item.get("title", "")).strip()
            for item in records
            if item.get("category") == "skill"
            and str(item.get("title", "")).strip()
        )
    )


def _render_education(
    document: Document,
    item: dict[str, Any],
) -> None:
    metadata = item.get("metadata", {})
    title_parts = [
        metadata.get("degree"),
        metadata.get("field_of_study"),
    ]

    title = " — ".join(
        str(part).strip()
        for part in title_parts
        if part
    )

    if not title:
        title = str(item.get("title") or "Education")

    institution = metadata.get("institution")
    details: list[str] = []

    if institution:
        details.append(str(institution))

    if metadata.get("cgpa") is not None:
        details.append(f"CGPA: {metadata['cgpa']}")

    start_date = metadata.get("start_date")
    end_date = metadata.get("end_date")

    if start_date or end_date:
        details.append(
            f"{start_date or ''} - {end_date or ''}".strip(" -")
        )

    if metadata.get("location"):
        details.append(str(metadata["location"]))

    _add_entry_heading(
        document,
        title,
        " | ".join(details) if details else None,
    )


def _render_experience(
    document: Document,
    item: dict[str, Any],
) -> None:
    metadata = item.get("metadata", {})
    role = metadata.get("role")
    company = metadata.get("company")

    title = " — ".join(
        str(value).strip()
        for value in (role, company)
        if value
    ) or item.get("title", "Experience")

    start_date = metadata.get("start_date")
    end_date = metadata.get("end_date")
    detail = (
        f"{start_date or ''} - {end_date or ''}".strip(" -")
        or None
    )

    if metadata.get("location"):
        detail = (
            f"{detail} | {metadata['location']}"
            if detail
            else str(metadata["location"])
        )

    _add_entry_heading(document, str(title), detail)

    bullets = metadata.get("bullets", [])

    if not isinstance(bullets, list):
        bullets = []

    for bullet in bullets:
        _add_bullet(document, str(bullet))

    if not bullets and item.get("statement"):
        _add_bullet(document, str(item["statement"]))


def _render_project(
    document: Document,
    item: dict[str, Any],
) -> None:
    _add_entry_heading(
        document,
        str(item.get("title") or "Project"),
    )

    metadata = item.get("metadata", {})
    bullets = metadata.get("bullets", [])

    if not isinstance(bullets, list):
        bullets = []

    if bullets:
        for bullet in bullets:
            _add_bullet(document, str(bullet))
    elif item.get("statement"):
        _add_bullet(document, str(item["statement"]))


def build_tailored_resume(
    selection: dict[str, Any],
    store: dict[str, Any],
    output_path: str | Path,
    *,
    audit_path: str | Path | None = None,
) -> Path:
    """Create an ATS-friendly DOCX using only currently verified evidence.

    Existing bullets are formatted, not rewritten. The builder rejects
    stale or unverified selection reports and does not invent a summary.
    """
    candidate = store.get("candidate")

    if not isinstance(candidate, dict) or not candidate.get("name"):
        raise ResumeBuildError("Evidence store has no candidate name.")

    records = _validated_records(selection, store)
    output = Path(output_path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    section = document.sections[0]

    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(2)

    bullet_style = document.styles["List Bullet"]
    bullet_style.font.name = "Arial"
    bullet_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    bullet_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    bullet_style.font.size = Pt(9.1)

    name_paragraph = document.add_paragraph()
    name_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_paragraph_spacing(name_paragraph, after=2)

    name_run = name_paragraph.add_run(str(candidate["name"]).strip())
    _set_run_font(
        name_run,
        size=19,
        bold=True,
        color=NAVY,
    )

    for contact_line in _contact_lines(candidate):
        _add_text(
            document,
            contact_line,
            size=8.3,
            color=MUTED,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            after=1,
        )

    by_category: dict[str, list[dict[str, Any]]] = {}

    for item in records:
        by_category.setdefault(
            str(item.get("category", "")), []
        ).append(item)

    # A one-column structure is used for straightforward ATS parsing.
    skill_text = _skills_text(by_category.get("skill", []))

    if skill_text:
        _add_section_heading(document, "Technical Skills")
        _add_text(
            document,
            skill_text,
            size=9.2,
            after=1,
        )

    experience = by_category.get("experience", [])

    if experience:
        _add_section_heading(document, "Experience")

        for item in experience:
            _render_experience(document, item)

    projects = by_category.get("project", [])

    if projects:
        _add_section_heading(document, "Projects")

        for item in projects:
            _render_project(document, item)

    education = by_category.get("education", [])

    if education:
        _add_section_heading(document, "Education")

        for item in education:
            _render_education(document, item)

    certifications = by_category.get("certification", [])

    if certifications:
        _add_section_heading(document, "Certifications")

        for item in certifications:
            _add_bullet(
                document,
                str(item.get("title") or item.get("statement", "")),
            )

    achievements = by_category.get("achievement", [])

    if achievements:
        _add_section_heading(document, "Achievements")

        for item in achievements:
            _add_bullet(document, str(item.get("statement", "")))

    document.core_properties.title = (
        f"Tailored Resume - "
        f"{selection.get('job', {}).get('title', 'Job Application')}"
    )
    document.core_properties.subject = (
        "Resume generated from user-verified evidence"
    )
    document.core_properties.author = "JobPilot"

    document.save(output)

    if audit_path is not None:
        audit = {
            "generated_at_utc": _utc_now(),
            "output_file": output.name,
            "job": selection.get("job", {}),
            "source_resume": store.get("source_resume", {}),
            "evidence_count": len(records),
            "evidence": [
                {
                    "evidence_id": item.get("evidence_id"),
                    "category": item.get("category"),
                    "title": item.get("title"),
                    "source": item.get("source", {}),
                    "review_status": item.get("review", {}).get("status"),
                    "user_verified": item.get("review", {}).get(
                        "user_verified"
                    ),
                    "included_as_core_section": bool(
                        item.get("included_as_core_section", False)
                    ),
                }
                for item in records
            ],
            "safety": {
                "only_currently_verified_evidence_used": True,
                "bullets_rewritten": False,
                "claims_generated": False,
            },
        }

        audit_file = Path(audit_path).expanduser().resolve()
        audit_file.parent.mkdir(parents=True, exist_ok=True)
        audit_file.write_text(
            json.dumps(audit, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )

    return output


def build_tailored_resume_from_files(
    selection_path: str | Path,
    store_path: str | Path,
    output_path: str | Path,
    *,
    audit_path: str | Path | None = None,
) -> Path:
    selection = _read_json(selection_path)
    store = _read_json(store_path)

    return build_tailored_resume(
        selection,
        store,
        output_path,
        audit_path=audit_path,
    )


def default_output_path(
    job_title: str,
    directory: str | Path,
) -> Path:
    return Path(directory).expanduser() / (
        f"Tailored_{_safe_slug(job_title)}.docx"
    )


__all__ = [
    "ResumeBuildError",
    "build_tailored_resume",
    "build_tailored_resume_from_files",
    "default_output_path",
]